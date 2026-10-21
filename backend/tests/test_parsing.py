from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pptx import Presentation

from app.core.errors import DomainError
from app.services import parsing
from app.services.chunking import build_chunks
from app.services.parsing import (
    ParsedBlock,
    ParsedDocument,
    detect_and_validate_file,
    normalize_extension,
    parse_document,
)


def test_markdown_preserves_headings_lists_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text(
        "# 蛋白质结构\n\n"
        "蛋白质是生命活动的重要承担者。\n\n"
        "## 一级结构\n\n"
        "- 氨基酸序列\n"
        "- 肽键连接\n\n"
        "| 层级 | 说明 |\n"
        "| --- | --- |\n"
        "| 一级 | 序列 |\n",
        encoding="utf-8",
    )

    parsed = parse_document(path, ".md", "notes.md")

    assert parsed.title == "蛋白质结构"
    assert [block.kind for block in parsed.blocks].count("heading") == 2
    assert any(block.kind == "list_item" for block in parsed.blocks)
    assert any(block.kind == "table" and "| 一级 |" in block.text for block in parsed.blocks)


def test_docx_preserves_heading_paragraph_and_table(tmp_path: Path) -> None:
    path = tmp_path / "chapter.docx"
    document = DocxDocument()
    document.add_heading("第二章 蛋白质", level=1)
    document.add_heading("蛋白质的分子结构", level=2)
    document.add_paragraph("蛋白质由氨基酸通过肽键连接而成。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "结构"
    table.cell(0, 1).text = "特征"
    table.cell(1, 0).text = "一级"
    table.cell(1, 1).text = "氨基酸序列"
    document.save(path)

    parsed = parse_document(path, ".docx", "chapter.docx")

    assert parsed.title == "第二章 蛋白质"
    assert any(
        block.kind == "heading" and block.level == 2 for block in parsed.blocks
    )
    assert any(block.kind == "table" and "氨基酸序列" in block.text for block in parsed.blocks)


def test_pptx_records_slide_numbers_and_text(tmp_path: Path) -> None:
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "蛋白质的二级结构"
    slide.placeholders[1].text = "α-螺旋\nβ-折叠"
    presentation.save(path)

    parsed = parse_document(path, ".pptx", "slides.pptx")

    assert parsed.slide_count == 1
    assert any(
        block.kind == "heading" and block.slide == 1 for block in parsed.blocks
    )
    assert any(
        block.slide == 1 and "β-折叠" in block.text for block in parsed.blocks
    )


def test_pdf_preserves_page_numbers(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "lecture.pdf"
    path.write_bytes(b"%PDF-1.4\n% test fixture\n")

    class FakePage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class FakeReader:
        is_encrypted = False

        def __init__(self, _path: str) -> None:
            self.pages = [
                FakePage("第一页：蛋白质由氨基酸组成。"),
                FakePage("第二页：肽键连接氨基酸。"),
            ]

    monkeypatch.setattr(parsing, "PdfReader", FakeReader)

    parsed = parse_document(path, ".pdf", "lecture.pdf")

    assert parsed.page_count == 2
    assert {block.page for block in parsed.blocks} == {1, 2}
    assert parsed.needs_ocr is False


def test_scanned_pdf_is_flagged_without_ocr(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "scan.pdf"
    path.write_bytes(b"%PDF-1.4\n% scanned fixture\n")

    class EmptyPage:
        def extract_text(self) -> str:
            return ""

    class EmptyReader:
        is_encrypted = False
        pages = [EmptyPage(), EmptyPage()]

    monkeypatch.setattr(parsing, "PdfReader", lambda _path: EmptyReader())

    parsed = parse_document(path, ".pdf", "scan.pdf")

    assert parsed.needs_ocr is True
    assert parsed.blocks == []
    assert len(parsed.warnings) == 2


def test_file_signature_validation_and_legacy_ppt(tmp_path: Path) -> None:
    path = tmp_path / "fake.pdf"
    path.write_text("not a pdf", encoding="utf-8")

    with pytest.raises(DomainError) as error:
        detect_and_validate_file(path, ".pdf")
    assert error.value.code == "file_content_mismatch"

    with pytest.raises(DomainError) as legacy:
        normalize_extension("old.ppt")
    assert legacy.value.code == "legacy_ppt_unsupported"


def test_chunking_preserves_heading_path_and_source_ranges() -> None:
    paragraph = "蛋白质的空间结构决定其功能。" * 40
    parsed = ParsedDocument(
        title="蛋白质",
        blocks=[
            ParsedBlock("heading", "蛋白质结构", level=1),
            ParsedBlock("paragraph", paragraph),
            ParsedBlock("heading", "功能关系", level=2),
            ParsedBlock("table", "| 结构 | 功能 |\n| --- | --- |\n| 一级 | 序列 |"),
        ],
        page_count=3,
    )

    chunks = build_chunks(parsed, max_chars=300, min_chars=20)

    assert len(chunks) >= 2
    assert chunks[0].heading_path == "蛋白质结构"
    assert any(chunk.heading_path == "蛋白质结构 > 功能关系" for chunk in chunks)
    assert all(chunk.char_end > chunk.char_start for chunk in chunks)
    assert all(chunk.content.strip() for chunk in chunks)
    assert any("| 结构 | 功能 |" in chunk.content for chunk in chunks)


def test_chunking_keeps_pdf_page_boundaries() -> None:
    parsed = ParsedDocument(
        title="分页资料",
        blocks=[
            ParsedBlock("paragraph", "第一页的蛋白质内容。", page=1),
            ParsedBlock("paragraph", "第二页的蛋白质内容。", page=2),
        ],
        page_count=2,
    )

    chunks = build_chunks(parsed, max_chars=6000, min_chars=10)

    assert len(chunks) == 2
    assert [chunk.page_start for chunk in chunks] == [1, 2]
    assert all(chunk.location_label.startswith("第 ") for chunk in chunks)


def test_chunking_reports_accurate_ranges_for_split_blocks() -> None:
    text = "A" * 200 + "。" + "B" * 200 + "。"
    parsed = ParsedDocument(
        title="长段落",
        blocks=[ParsedBlock("paragraph", text, page=1)],
        page_count=1,
    )

    chunks = build_chunks(parsed, max_chars=300, min_chars=0)

    assert len(chunks) == 2
    assert [(chunk.char_start, chunk.char_end) for chunk in chunks] == [
        (0, 201),
        (201, 402),
    ]
    assert chunks[0].content == text[0:201]
    assert chunks[1].content == text[201:402]
