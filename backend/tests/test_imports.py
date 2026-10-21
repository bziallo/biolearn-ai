from __future__ import annotations

from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.core.config import Settings
from app.models import (
    Chapter,
    Document,
    ImportJob,
    KnowledgeExtractionJob,
    KnowledgePoint,
)
from app.services import documents as document_service
from app.services import importing as import_service
from app.services import knowledge as knowledge_service


class FakeAI:
    def __init__(self, responses: list[dict]) -> None:
        self.responses = list(responses)
        self.calls: list[tuple[str, str]] = []

    def chat_json(self, system_prompt: str, user_prompt: str, *, cancel_check=None) -> dict:
        self.calls.append((system_prompt, user_prompt))
        if not self.responses:
            raise AssertionError("FakeAI ran out of responses")
        return self.responses.pop(0)


ANALYSIS = {
    "theme": "蛋白质的结构与功能",
    "learning_objectives": ["理解一级结构"],
    "sections": ["一级结构", "二级结构"],
    "core_concepts": ["氨基酸序列"],
    "processes": ["肽键形成"],
    "relationships": ["序列决定结构"],
    "formulas": [],
    "confusions": [],
    "prerequisites": ["氨基酸"],
    "cross_links": [],
}

EXTRACTION = {
    "points": [
        {
            "name": "蛋白质一级结构",
            "type": "structure",
            "explanation": "蛋白质的一级结构是氨基酸的线性排列顺序。",
            "core_points": ["氨基酸序列"],
            "conditions": None,
            "source_quote": "蛋白质由氨基酸通过肽键连接。",
            "confidence": 0.9,
            "needs_review": False,
            "parent_name": None,
        }
    ]
}


@pytest.fixture
def import_env(monkeypatch: pytest.MonkeyPatch, app_settings: Settings):
    real_run = import_service.run_import
    real_run_analysis = import_service.run_import_analysis
    monkeypatch.setattr(import_service, "get_settings", lambda: app_settings)
    monkeypatch.setattr(document_service, "get_settings", lambda: app_settings)
    monkeypatch.setattr(import_service, "run_import", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        import_service,
        "run_import_analysis",
        lambda *_args, **_kwargs: None,
    )
    return real_run, real_run_analysis


def create_subject(client: TestClient, name: str) -> dict:
    response = client.post("/api/v1/subjects", json={"name": name})
    assert response.status_code == 201, response.text
    return response.json()


def create_chapter(client: TestClient, subject_id: str, name: str) -> dict:
    response = client.post(
        f"/api/v1/subjects/{subject_id}/chapters",
        json={"name": name},
    )
    assert response.status_code == 201, response.text
    return response.json()


def make_pptx(
    path: Path,
    *,
    title: str = "蛋白质的结构与功能",
    toc: list[str] | None = None,
    slides: list[tuple[str, str]] | None = None,
) -> Path:
    presentation = Presentation()
    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = title
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = "生物化学"

    if toc:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "目录"
        slide.placeholders[1].text = "\n".join(toc)

    for slide_title, body in slides or []:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_title
        slide.placeholders[1].text = body

    presentation.save(path)
    return path


def upload_import(
    client: TestClient,
    path: Path,
    *,
    subject_id: str | None,
    mode: str = "auto_create",
    chapter_id: str | None = None,
    auto_analyze: bool = False,
    require_preview: bool = False,
) -> dict:
    data: dict[str, str] = {
        "mode": mode,
        "auto_analyze": "true" if auto_analyze else "false",
        "require_preview": "true" if require_preview else "false",
    }
    if subject_id:
        data["subject_id"] = subject_id
    if chapter_id:
        data["chapter_id"] = chapter_id
    response = client.post(
        "/api/v1/imports",
        data=data,
        files={
            "file": (
                path.name,
                path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def chapter_names(client: TestClient, subject_id: str) -> list[str]:
    response = client.get(f"/api/v1/subjects/{subject_id}/chapters")
    assert response.status_code == 200
    tree = response.json()
    names: list[str] = []

    def walk(nodes: list[dict]) -> None:
        for node in nodes:
            names.append(node["name"])
            walk(node.get("children", []))

    walk(tree)
    return names


def test_import_into_preset_chapter_without_creating_new_chapters(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    chapter = create_chapter(client, subject["id"], "蛋白质")
    path = make_pptx(
        tmp_path / "preset.pptx",
        toc=["第一章 蛋白质的分子结构", "第一节 一级结构"],
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )

    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="preset",
        chapter_id=chapter["id"],
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )

    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert result["status"] == "completed"
    assert result["result"]["chapter_id"] == chapter["id"]
    assert chapter_names(client, subject["id"]) == ["蛋白质"]

    detail = client.get(f"/api/v1/documents/{result['document_id']}").json()
    assert detail["chapter_name"] == "蛋白质"
    with session_factory() as db:
        assert (
            db.query(KnowledgeExtractionJob)
            .filter(KnowledgeExtractionJob.document_id == result["document_id"])
            .count()
            == 0
        )


def test_import_auto_creates_ppt_outline_with_slide_sources(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    path = make_pptx(
        tmp_path / "auto.pptx",
        toc=[
            "第一章 蛋白质的分子结构",
            "第一节 一级结构",
            "第二节 二级结构",
        ],
        slides=[
            ("一级结构", "蛋白质由氨基酸通过肽键连接。"),
            ("二级结构", "α-螺旋和β-折叠是常见二级结构。"),
        ],
    )

    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="auto_create",
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )

    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert result["status"] == "completed"
    assert result["result"]["created_chapter_ids"]
    names = chapter_names(client, subject["id"])
    assert names[0] == "蛋白质的结构与功能"
    assert "第一章 蛋白质的分子结构" in names
    assert "第一节 一级结构" in names
    assert "第二节 二级结构" in names

    with session_factory() as db:
        created = (
            db.query(Chapter)
            .filter(Chapter.id.in_(result["result"]["created_chapter_ids"]))
            .all()
        )
        assert all(chapter.source_document_id for chapter in created)
        assert all(chapter.source_import_job_id for chapter in created)
        assert any(chapter.source_slide is not None for chapter in created)


def test_import_preview_confirm_then_archive(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    path = make_pptx(
        tmp_path / "preview.pptx",
        toc=["第一章 蛋白质的分子结构", "第一节 一级结构"],
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )

    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="auto_create",
        require_preview=True,
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )

    pending = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert pending["status"] == "awaiting_confirmation"
    assert pending["suggestion"]["chapters"]
    assert chapter_names(client, subject["id"]) == []

    response = client.post(
        f"/api/v1/imports/{imported['id']}/confirm",
        json={"accept_suggestions": True, "auto_analyze": False},
    )
    assert response.status_code == 200, response.text
    confirmed = response.json()
    assert confirmed["status"] == "completed"
    assert confirmed["result"]["chapter_id"]
    assert "第一章 蛋白质的分子结构" in chapter_names(client, subject["id"])


def test_import_same_file_does_not_create_duplicate_chapters(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    path = make_pptx(
        tmp_path / "duplicate.pptx",
        toc=["第一章 蛋白质的分子结构", "第一节 一级结构"],
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )

    first = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="auto_create",
    )
    real_run(
        first["id"],
        session_factory=session_factory,
        settings=app_settings,
    )
    before = chapter_names(client, subject["id"])

    second = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="auto_create",
    )
    after = client.get(f"/api/v1/imports/{second['id']}").json()
    assert after["status"] == "completed"
    assert after["result"]["duplicate"] is True
    assert chapter_names(client, subject["id"]) == before


def test_import_auto_analyze_and_cancel_keeps_document_and_chapters(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, real_run_analysis = import_env
    subject = create_subject(client, "生物化学")
    path = make_pptx(
        tmp_path / "analysis.pptx",
        toc=["第一章 蛋白质的分子结构", "第一节 一级结构"],
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )

    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="auto_create",
        auto_analyze=True,
    )
    fake = FakeAI([ANALYSIS, EXTRACTION])
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
        client=fake,
    )

    started = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert started["status"] == "analyzing"
    assert started["analysis_job_id"]

    # Cancelling a queued analysis must not remove the imported document or
    # the chapters created by the import.
    with session_factory() as db:
        job = knowledge_service.cancel_extraction_job(db, started["analysis_job_id"])
        assert job.status == "cancelled"

    real_run_analysis(
        imported["id"],
        started["analysis_job_id"],
        session_factory=session_factory,
        settings=app_settings,
        client=fake,
    )

    after = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert after["status"] == "completed"
    assert after["analysis_status"] == "cancelled"
    assert after["result"]["chapter_id"]

    document = client.get(f"/api/v1/documents/{after['document_id']}").json()
    assert document["status"] == "parsed"
    assert document["chapter_id"]
    with session_factory() as db:
        assert db.get(Document, after["document_id"]) is not None
        assert db.get(Chapter, after["result"]["chapter_id"]) is not None


def test_import_without_default_subject_requires_selection(
    client: TestClient,
    import_env,
    tmp_path: Path,
) -> None:
    path = make_pptx(
        tmp_path / "no-subject.pptx",
        slides=[("内容", "蛋白质由氨基酸通过肽键连接。")],
    )
    response = client.post(
        "/api/v1/imports",
        data={"mode": "auto_create", "auto_analyze": "false"},
        files={"file": (path.name, path.read_bytes(), "application/octet-stream")},
    )
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "default_subject_missing"


def test_import_preferences_are_persisted_and_used_as_default(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    chapter = create_chapter(client, subject["id"], "蛋白质")

    response = client.put(
        "/api/v1/imports/preferences",
        json={
            "default_subject_id": subject["id"],
            "default_chapter_id": chapter["id"],
            "default_mode": "preset",
            "auto_analyze": False,
            "require_preview": False,
            "uncertain_policy": "confirm",
        },
    )
    assert response.status_code == 200, response.text
    preferences = client.get("/api/v1/imports/preferences").json()
    assert preferences["default_subject_id"] == subject["id"]
    assert preferences["default_chapter_id"] == chapter["id"]
    assert preferences["default_mode"] == "preset"

    path = make_pptx(
        tmp_path / "defaults.pptx",
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )
    imported = upload_import(
        client,
        path,
        subject_id=None,
        mode="preset",
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )
    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert result["status"] == "completed"
    assert result["result"]["chapter_id"] == chapter["id"]


def test_import_without_toc_uses_conservative_single_chapter(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    path = make_pptx(
        tmp_path / "generic.pptx",
        slides=[
            ("概述", "这一页没有编号，也没有目录。"),
            ("补充", "另一页普通内容，不应被当成章节标题。"),
        ],
    )
    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="auto_create",
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )
    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert result["status"] == "completed"
    assert len(chapter_names(client, subject["id"])) == 1
    assert chapter_names(client, subject["id"])[0] == "蛋白质的结构与功能"


def test_import_chapter_creation_failure_rolls_back_all_chapters(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    path = make_pptx(
        tmp_path / "failure.pptx",
        toc=["第一章 蛋白质的分子结构", "第一节 一级结构", "第二节 二级结构"],
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )

    calls = {"count": 0}
    original = import_service._next_sort_order

    def failing_next_sort_order(*args, **kwargs):
        calls["count"] += 1
        if calls["count"] >= 2:
            raise RuntimeError("simulated chapter failure")
        return original(*args, **kwargs)

    monkeypatch.setattr(
        import_service,
        "_next_sort_order",
        failing_next_sort_order,
    )

    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="auto_create",
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )

    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert result["status"] == "failed"
    assert result["error"]
    assert chapter_names(client, subject["id"]) == []
    document = client.get(f"/api/v1/documents/{result['document_id']}").json()
    assert document["status"] == "parsed"


def test_import_auto_analyze_completes_and_creates_drafts(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, real_run_analysis = import_env
    subject = create_subject(client, "生物化学")
    path = tmp_path / "analysis-ok.txt"
    path.write_text("蛋白质由氨基酸通过肽键连接。", encoding="utf-8")
    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="auto_create",
        auto_analyze=True,
    )
    fake = FakeAI([ANALYSIS, EXTRACTION])
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
        client=fake,
    )
    started = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert started["status"] == "analyzing"

    real_run_analysis(
        imported["id"],
        started["analysis_job_id"],
        session_factory=session_factory,
        settings=app_settings,
        client=fake,
    )
    finished = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert finished["status"] == "completed"
    assert finished["analysis_status"] == "completed"
    with session_factory() as db:
        points = (
            db.query(KnowledgePoint)
            .filter(KnowledgePoint.document_id == finished["document_id"])
            .all()
        )
        assert len(points) == 1
        assert points[0].name == "蛋白质一级结构"


def test_document_attribution_can_be_adjusted_after_import(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject_a = create_subject(client, "生物化学")
    subject_b = create_subject(client, "分子生物学")
    chapter_b = create_chapter(client, subject_b["id"], "蛋白质结构")
    path = make_pptx(
        tmp_path / "move-document.pptx",
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )
    imported = upload_import(
        client,
        path,
        subject_id=subject_a["id"],
        mode="auto_create",
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )
    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    document_id = result["document_id"]

    point = client.post(
        "/api/v1/knowledge-points",
        json={
            "document_id": document_id,
            "name": "手动知识点",
            "explanation": "用于归属调整测试。",
        },
    )
    assert point.status_code == 201, point.text

    response = client.patch(
        f"/api/v1/documents/{document_id}",
        json={"subject_id": subject_b["id"], "chapter_id": chapter_b["id"]},
    )
    assert response.status_code == 200, response.text
    moved = response.json()
    assert moved["subject_id"] == subject_b["id"]
    assert moved["chapter_id"] == chapter_b["id"]

    point_detail = client.get(
        f"/api/v1/knowledge-points/{point.json()['id']}"
    ).json()
    assert point_detail["subject_id"] == subject_b["id"]
    assert point_detail["chapter_id"] == chapter_b["id"]


def test_smart_archive_reuses_matching_root_and_creates_missing_child(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    root = create_chapter(client, subject["id"], "蛋白质的结构与功能")
    path = make_pptx(
        tmp_path / "smart-archive.pptx",
        toc=["第一章 蛋白质的分子结构", "第一节 一级结构"],
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )
    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="smart_archive",
        require_preview=False,
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )
    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert result["status"] == "completed"
    assert root["id"] in result["result"]["reused_chapter_ids"]
    assert result["result"]["created_chapter_ids"]
    assert result["result"]["chapter_id"] == root["id"]
    names = chapter_names(client, subject["id"])
    assert names.count("蛋白质的结构与功能") == 1
    assert "第一章 蛋白质的分子结构" in names


def test_smart_archive_similar_name_requires_confirmation(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    create_chapter(client, subject["id"], "蛋白质结构与功能")
    path = make_pptx(
        tmp_path / "similar.pptx",
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )
    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="smart_archive",
        require_preview=False,
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )
    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert result["status"] == "awaiting_confirmation"
    assert result["suggestion"]["chapters"][0]["action"] == "uncertain"
    assert chapter_names(client, subject["id"]) == ["蛋白质结构与功能"]


def test_smart_archive_exact_match_skips_preview_when_nothing_is_created(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject = create_subject(client, "生物化学")
    root = create_chapter(client, subject["id"], "蛋白质的结构与功能")
    path = make_pptx(
        tmp_path / "exact-match.pptx",
        slides=[("概述", "这一页没有编号，也没有目录。")],
    )
    imported = upload_import(
        client,
        path,
        subject_id=subject["id"],
        mode="smart_archive",
        require_preview=True,
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )
    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert result["status"] == "completed"
    assert result["result"]["chapter_id"] == root["id"]
    assert result["result"]["created_chapter_ids"] == []
    assert chapter_names(client, subject["id"]) == ["蛋白质的结构与功能"]


def test_same_chapter_name_in_different_subjects_is_not_reused(
    client: TestClient,
    session_factory,
    app_settings: Settings,
    import_env,
    tmp_path: Path,
) -> None:
    real_run, _ = import_env
    subject_a = create_subject(client, "生物化学")
    subject_b = create_subject(client, "分子生物学")
    chapter_a = create_chapter(client, subject_a["id"], "蛋白质的结构与功能")

    path = make_pptx(
        tmp_path / "different-subject.pptx",
        toc=["第一章 蛋白质的分子结构"],
        slides=[("一级结构", "蛋白质由氨基酸通过肽键连接。")],
    )
    imported = upload_import(
        client,
        path,
        subject_id=subject_b["id"],
        mode="auto_create",
    )
    real_run(
        imported["id"],
        session_factory=session_factory,
        settings=app_settings,
    )
    result = client.get(f"/api/v1/imports/{imported['id']}").json()
    assert result["status"] == "completed"
    assert result["result"]["chapter_id"] != chapter_a["id"]
    assert result["result"]["subject_id"] == subject_b["id"]
    assert chapter_names(client, subject_b["id"])[0] == "蛋白质的结构与功能"
    with session_factory() as db:
        assert db.get(Chapter, result["result"]["chapter_id"]).subject_id == subject_b["id"]
